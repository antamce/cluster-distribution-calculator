from __future__ import annotations

import contextlib
import shutil
import unittest
import uuid
from pathlib import Path

from openpyxl import load_workbook

from synpo.exporting import (
    _write_workbook,
    collect_export_tables_from_results,
    filter_exported_measurement_workbook,
)
from synpo.measurements import (
    apply_spine_volume_filter,
    spine_volume_distribution_rows,
)


@contextlib.contextmanager
def workspace_directory():
    root = Path.cwd() / ".test-work" / uuid.uuid4().hex
    root.mkdir(parents=True)
    try:
        yield root
    finally:
        shutil.rmtree(root)


def synthetic_result() -> dict[str, object]:
    common = {"experimental_group": "treated", "specimen_id": "cell-1", "roi_id": 1, "dendrite_id": 1}
    spines = [
        {**common, "spine_id": 1, "volume_um3": 0.10, "spine_valid": True, "has_protein_cluster": True, "included_cluster_count": 1, "cluster_to_spine_volume_ratio": 0.2},
        {**common, "spine_id": 2, "volume_um3": 0.20, "spine_valid": True, "has_protein_cluster": False, "included_cluster_count": 0, "cluster_to_spine_volume_ratio": None},
        {**common, "spine_id": 3, "volume_um3": 0.05, "spine_valid": False, "validity_reviewed": True, "validity_note": "manual", "has_protein_cluster": False, "included_cluster_count": 0, "cluster_to_spine_volume_ratio": None},
        {**common, "spine_id": 4, "volume_um3": 0.15, "spine_valid": True, "has_protein_cluster": True, "included_cluster_count": 1, "cluster_to_spine_volume_ratio": 0.1},
    ]
    clusters = [
        {**common, "spine_id": spine_id, "cluster_id": spine_id, "row_type": "individual_cluster", "volume_inside_spine_um3": volume, "spine_valid": True}
        for spine_id, volume in ((1, 0.02), (4, 0.015))
    ] + [
        {**common, "spine_id": spine_id, "row_type": "spine_cluster_sum", "included_cluster_count": 1, "spine_valid": True}
        for spine_id in (1, 4)
    ]
    distributions = [
        {
            **common,
            "spine_id": spine_id,
            "spine_valid": spine_id != 3,
            "distribution_included": True,
            **{f"bin_{index:02d}_ratio": spine_id / 10.0 for index in range(1, 11)},
        }
        for spine_id in (1, 4)
    ]
    return {
        "spine_rows": spines,
        "cluster_rows": clusters,
        "distribution_rows": distributions,
        "dendrite_rows": [{**common, "length_um": 10.0}],
        "roi_rows": [{"experimental_group": "treated", "specimen_id": "cell-1", "roi_id": 1}],
        "specimen_rows": [{"experimental_group": "treated", "specimen_id": "cell-1", "dendrite_count": 1}],
    }


class SpineVolumeFilterTests(unittest.TestCase):
    def test_strict_cutoff_override_and_manual_invalid_are_independent(self) -> None:
        filtered, audit = apply_spine_volume_filter(
            synthetic_result(), enabled=True, cutoff_um3=0.20, force_keep_ids={4}
        )
        by_id = {int(row["spine_id"]): row for row in filtered["spine_rows"]}
        self.assertTrue(by_id[2]["spine_valid"], "equality with cutoff must be retained")
        self.assertTrue(by_id[4]["spine_valid"], "force keep must override volume only")
        self.assertFalse(by_id[3]["spine_valid"], "manual invalid must remain invalid")
        self.assertEqual([int(row["spine_id"]) for row in audit], [1])
        self.assertEqual(filtered["specimen_rows"][0]["spine_count"], 2)
        self.assertEqual(filtered["specimen_rows"][0]["invalid_spine_count"], 1)
        self.assertEqual(filtered["specimen_rows"][0]["volume_filtered_spine_count"], 1)
        self.assertEqual(
            {int(row["spine_id"]) for row in filtered["cluster_rows"] if row["spine_valid"]},
            {4},
        )

    def test_shared_volume_distribution_contains_pooled_and_weighted_rows(self) -> None:
        filtered, _ = apply_spine_volume_filter(
            synthetic_result(), enabled=True, cutoff_um3=0.20, force_keep_ids={4}
        )
        rows = spine_volume_distribution_rows(
            [filtered], filter_enabled=True, cutoff_um3=0.20
        )
        metadata = rows[0]
        self.assertEqual(metadata["valid_spine_count"], 2)
        self.assertIn("bin_edges_um3", metadata)
        descriptive = [row for row in rows if row.get("row_type") == "descriptive"]
        self.assertEqual({row["scope"] for row in descriptive}, {"All", "treated"})
        bins = [row for row in rows if row.get("row_type") == "histogram_bin"]
        self.assertTrue(bins)
        self.assertTrue(all(row["n_specimens"] == 1 for row in bins))

    def test_standalone_filter_preserves_custom_sheet_and_recalculates(self) -> None:
        with workspace_directory() as root:
            manifest = {
                "measurements": {"settings": {"spine_volume_filter_enabled": False, "spine_volume_filter_cutoff_um3": 0.0}},
                "calibration": {},
                "specimens": [],
            }
            tables = collect_export_tables_from_results(manifest, [synthetic_result()])
            source = root / "source.xlsx"
            output = root / "source_volume-filtered.xlsx"
            _write_workbook(source, tables)
            workbook = load_workbook(source)
            custom = workbook.create_sheet("User Notes")
            custom["A1"] = "preserve me"
            workbook.save(source)
            workbook.close()

            result = filter_exported_measurement_workbook(
                source,
                output,
                cutoff_um3=0.20,
                force_keep_keys={("treated", "cell-1", 4)},
            )
            self.assertTrue(result["verified"])
            self.assertEqual(result["excluded_spine_count"], 1)
            self.assertEqual(
                [
                    row["table"]
                    for row in result["compatibility_report"]
                    if row["status"] == "unavailable" and row["table"] != "PDF exports"
                ],
                [],
            )
            self.assertTrue(source.is_file())
            self.assertTrue((root / "source_volume-filtered_csv" / "Compatibility_Report.csv").is_file())
            filtered = load_workbook(output, read_only=True, data_only=True)
            try:
                self.assertEqual(filtered["User Notes"]["A1"].value, "preserve me")
                spine_rows = list(filtered["Spine_Master"].iter_rows(values_only=True))
                spine_header = list(spine_rows[0])
                spine_id_index = spine_header.index("spine_id")
                self.assertEqual({row[spine_id_index] for row in spine_rows[1:]}, {2, 3, 4})
                dendrite_rows = list(filtered["Dendrite_Master"].iter_rows(values_only=True))
                dendrite_header = list(dendrite_rows[0])
                self.assertEqual(dendrite_rows[1][dendrite_header.index("spine_count")], 2)
                audit_rows = list(filtered["Volume_Filtered_Spines"].iter_rows(values_only=True))
                self.assertEqual(len(audit_rows), 2)
                report_rows = list(filtered["Compatibility_Report"].iter_rows(values_only=True))
                self.assertTrue(any("recalculated" in row for row in report_rows[1:]))
            finally:
                filtered.close()
            with self.assertRaisesRegex(ValueError, "original unfiltered workbook"):
                filter_exported_measurement_workbook(
                    output,
                    root / "second-pass.xlsx",
                    cutoff_um3=0.1,
                )

    def test_standalone_legacy_missing_columns_are_reported_not_approximated(self) -> None:
        with workspace_directory() as root:
            manifest = {
                "measurements": {"settings": {}},
                "calibration": {},
                "specimens": [],
            }
            source = root / "legacy.xlsx"
            output = root / "legacy_volume-filtered.xlsx"
            _write_workbook(
                source,
                collect_export_tables_from_results(manifest, [synthetic_result()]),
            )
            workbook = load_workbook(source)
            sheet = workbook["Distribution_Individual"]
            headers = [cell.value for cell in sheet[1]]
            column = headers.index("bin_10_ratio") + 1
            sheet.delete_cols(column)
            workbook.save(source)
            workbook.close()

            result = filter_exported_measurement_workbook(
                source, output, cutoff_um3=0.2
            )
            unavailable = {
                row["table"]: row["detail"]
                for row in result["compatibility_report"]
                if row["status"] == "unavailable"
            }
            self.assertIn("Distribution_Group", unavailable)
            self.assertIn("Dendrite_Master", unavailable)
            filtered = load_workbook(output, read_only=True, data_only=True)
            try:
                self.assertEqual(
                    filtered["Distribution_Group"]["A2"].value,
                    "NOT RECALCULATED",
                )
                self.assertIn("bin_10_ratio", unavailable["Distribution_Group"])
            finally:
                filtered.close()


if __name__ == "__main__":
    unittest.main()
