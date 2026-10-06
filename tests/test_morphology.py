from __future__ import annotations

from pathlib import Path
from contextlib import contextmanager
import shutil
import uuid

import pytest
from openpyxl import load_workbook
from matplotlib.figure import Figure
import numpy as np

from synpo.morphology import (
    DEFAULT_FEATURES,
    MorphologyClusteringSettings,
    draw_pca_3d_feature_axes,
    draw_pca_interpretation,
    export_morphology_analysis,
    load_morphology_rows_from_workbook,
    morphology_feature_value,
    pca_interpretation_data,
    run_morphology_clustering,
    run_morphology_clustering_from_workbook,
)


@contextmanager
def workspace_directory():
    root = Path.cwd() / ".test-work" / uuid.uuid4().hex
    root.mkdir(parents=True)
    try:
        yield root
    finally:
        shutil.rmtree(root, ignore_errors=True)


def morphology_rows() -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for index in range(80):
        large = index >= 40
        offset = (index % 10) * 0.002
        rows.append(
            {
                "experimental_group": "treated" if index % 2 else "control",
                "specimen_id": f"specimen-{index // 10:02d}",
                "roi_id": 1,
                "dendrite_id": index // 5 + 1,
                "spine_id": index + 1,
                "spine_valid": True,
                "volume_um3": (0.12 if large else 0.025) + offset,
                "spine_curvilinear_length_um": (1.8 if large else 0.65) + offset,
                "spine_base_to_tip_distance_um": (1.6 if large else 0.58) + offset,
                "centerline_tortuosity": (1.12 if large else 1.05) + offset,
                "maximum_width_um": (0.7 if large else 0.28) + offset,
                "principal_axis_elongation": (2.1 if large else 1.3) + offset,
                "head_maximum_width_um": (0.62 if large else 0.24) + offset,
                "neck_median_width_um": (0.18 if large else 0.12) + offset,
                "head_to_neck_width_ratio": 3.4 if large else 2.0,
                "head_volume_um3": (0.09 if large else 0.015) + offset,
                "neck_volume_um3": (0.03 if large else 0.01) + offset,
                "surface_area_um2": (1.2 if large else 0.4) + offset,
                "sphericity": 0.72 if large else 0.61,
                "has_protein_cluster": index % 3 == 0,
                "protein_puncta_count": index % 4,
                "protein_puncta_total_volume_um3": index * 0.001,
                "protein_puncta_to_spine_volume_ratio": index * 0.002,
                "protein_puncta_mean_volume_um3": index * 0.0005,
                "protein_puncta_axis_centroid_fraction": (index % 10) / 10,
                "protein_distribution_in_spine": [float(bin_index + index % 2) for bin_index in range(10)],
            }
        )
    return rows


@pytest.mark.parametrize("algorithm", ["gaussian_mixture", "ward", "kmeans"])
def test_all_three_algorithms_and_pca_dimensions(algorithm: str) -> None:
    result = run_morphology_clustering(
        morphology_rows(),
        MorphologyClusteringSettings(
            algorithm=algorithm,
            features=DEFAULT_FEATURES,
            pca_dimensions=3,
            fixed_cluster_count=2,
        ),
    )
    assert result["selected_cluster_count"] == 2
    assert result["included_spine_count"] == 80
    assert result["protein_metrics_used_for_clustering"] is False
    assert all("pca_3" in row for row in result["assignments"])
    assert all(
        "clustering_feature_log_volume" in row
        for row in result["assignments"]
    )
    assert {row["subset"] for row in result["protein_summary"]} == {
        "all_spines",
        "protein_positive",
    }


def test_protein_values_do_not_change_clusters() -> None:
    settings = MorphologyClusteringSettings(fixed_cluster_count=2)
    original = morphology_rows()
    changed = [{**row, "protein_puncta_count": 10_000 - index} for index, row in enumerate(original)]
    first = run_morphology_clustering(original, settings)
    second = run_morphology_clustering(changed, settings)
    assert [row["morphology_cluster_id"] for row in first["assignments"]] == [
        row["morphology_cluster_id"] for row in second["assignments"]
    ]


@pytest.mark.parametrize("selection", ["silhouette", "elbow"])
def test_automatic_cluster_count_can_use_silhouette_or_elbow(selection: str) -> None:
    result = run_morphology_clustering(
        morphology_rows(),
        MorphologyClusteringSettings(
            algorithm="kmeans",
            features=DEFAULT_FEATURES,
            minimum_clusters=1,
            maximum_clusters=5,
            fixed_cluster_count=0,
            cluster_count_selection=selection,
            minimum_cluster_spines=5,
            minimum_cluster_fraction=0.0,
        ),
    )
    assert result["selected_cluster_count"] == 2
    selected = [
        row for row in result["candidate_diagnostics"] if row["selected"]
    ]
    assert len(selected) == 1
    assert selected[0]["cluster_count"] == 2
    assert all("within_cluster_sse" in row for row in result["candidate_diagnostics"])
    if selection == "elbow":
        assert selected[0]["elbow_distance"] is not None


def test_pca_interpretation_equations_reproduce_exported_scores() -> None:
    result = run_morphology_clustering(
        morphology_rows(),
        MorphologyClusteringSettings(fixed_cluster_count=2, pca_dimensions=3),
    )
    interpretation = pca_interpretation_data(result)
    assert len(interpretation["components"]) == 3
    assert "log10(volume_um3)" in interpretation["components"][0][
        "expanded_equation"
    ]
    assignment = result["assignments"][0]
    for component_index, component in enumerate(
        interpretation["components"], start=1
    ):
        score = 0.0
        for feature_row in interpretation["feature_rows"]:
            value = morphology_feature_value(
                assignment, str(feature_row["feature"])
            )
            score += float(feature_row[f"pca_{component_index}_loading"]) * (
                (float(value) - float(feature_row["scaling_center"]))
                / float(feature_row["scaling_scale"])
                - float(feature_row["pca_center_after_scaling"])
            )
        assert score == pytest.approx(assignment[f"pca_{component_index}"], abs=1e-8)
    figure = Figure(figsize=(10, 7))
    draw_pca_interpretation(figure, result)
    score_axis = figure.axes[0]
    assert score_axis.name == "3d"
    assert score_axis.get_zlabel().startswith("PC3")
    standalone_figure = Figure(figsize=(10, 7))
    draw_pca_3d_feature_axes(standalone_figure, result)
    standalone_axis = standalone_figure.axes[0]
    assert standalone_axis.name == "3d"
    assert standalone_axis.get_title().startswith("3D PCA scores")
    visible_collection_count = len(standalone_axis.collections)
    result["plot_style"] = {
        "pca_show_points": False,
        "pca_point_alpha": 0.1,
    }
    vectors_only_figure = Figure(figsize=(10, 7))
    draw_pca_3d_feature_axes(vectors_only_figure, result)
    assert len(vectors_only_figure.axes[0].collections) < visible_collection_count


def test_clustering_can_be_limited_to_selected_experimental_groups() -> None:
    rows = morphology_rows()
    for index, source in enumerate(rows[:40]):
        rows.append({**source, "spine_id": 1000 + index})
    result = run_morphology_clustering(
        rows,
        MorphologyClusteringSettings(
            features=("log_volume",),
            fixed_cluster_count=2,
            included_groups=("control",),
        ),
    )
    assert result["included_spine_count"] == 60
    assert {row["experimental_group"] for row in result["assignments"]} == {"control"}
    assert any(
        row["clustering_exclusion_reason"]
        == "experimental group not selected for this run"
        for row in result["excluded_audit"]
    )


def test_clustering_can_be_limited_to_geometry_reviewed_spines() -> None:
    rows = morphology_rows()
    for index, row in enumerate(rows):
        row["geometry_reviewed"] = index < 60
    result = run_morphology_clustering(
        rows,
        MorphologyClusteringSettings(
            features=("log_volume",),
            fixed_cluster_count=2,
            reviewed_only=True,
        ),
    )
    assert result["included_spine_count"] == 60
    assert result["excluded_spine_count"] == 20
    assert result["settings"]["reviewed_only"] is True
    assert {
        row["clustering_exclusion_reason"] for row in result["excluded_audit"]
    } == {"geometry review not completed"}


def test_nonlinear_embedding_is_the_actual_clustering_space(monkeypatch) -> None:
    def fake_embedding(values, settings, *, seed=None):
        embedded = np.zeros((len(values), settings.embedding_dimensions), dtype=float)
        embedded[1::2, 0] = 100.0
        embedded[:, 1] = np.arange(len(values), dtype=float) * 0.0001
        return embedded, {
            "method": settings.reduction_method,
            "dimensions": settings.embedding_dimensions,
            "requested_neighbors": settings.umap_n_neighbors,
            "effective_neighbors": settings.umap_n_neighbors,
            "implementation": "test reducer",
            "implementation_version": "1",
        }

    monkeypatch.setattr("synpo.morphology._fit_nonlinear_embedding", fake_embedding)
    monkeypatch.setattr(
        "synpo.morphology._embedding_quality",
        lambda source, embedded, neighbors: {
            "trustworthiness": 1.0,
            "distance_rank_correlation": 0.5,
        },
    )
    result = run_morphology_clustering(
        morphology_rows(),
        MorphologyClusteringSettings(
            algorithm="kmeans",
            features=DEFAULT_FEATURES,
            fixed_cluster_count=2,
            reduction_method="umap",
            embedding_dimensions=5,
            embedding_plot_dimensions=3,
            embedding_stability_repetitions=1,
        ),
    )
    even_labels = {
        row["morphology_cluster_id"] for row in result["assignments"][0::2]
    }
    odd_labels = {
        row["morphology_cluster_id"] for row in result["assignments"][1::2]
    }
    assert len(even_labels) == len(odd_labels) == 1
    assert even_labels != odd_labels
    assert all("embedding_5" in row for row in result["assignments"])
    assert result["embedding_metadata"]["used_for_clustering"] is True
    assert result["settings"]["reduction_method"] == "umap"


def test_export_and_standalone_workbook_round_trip() -> None:
    with workspace_directory() as root:
        settings = MorphologyClusteringSettings(
            fixed_cluster_count=2, pca_dimensions=3
        )
        result = run_morphology_clustering(morphology_rows(), settings)
        result.update(
            {
                "name": "test run",
                "run_id": "test-run",
                "plot_style": {
                    "axes_color": "#123456",
                    "axes_alpha": 0.65,
                    "background_color": "#f0e0d0",
                    "background_alpha": 0.4,
                    "show_legend": False,
                    "legend_position": "outside_bottom",
                    "custom_x_feature": "maximum_width",
                    "custom_y_feature": "elongation",
                },
            }
        )
        output = export_morphology_analysis(result, root / "analysis.xlsx")
        assert Path(output["workbook"]).is_file()
        assert Path(output["report_pdf"]).is_file()
        assert (Path(output["plot_directory"]) / "volume_vs_length.svg").is_file()
        assert (
            Path(output["plot_directory"])
            / "morphology_maximum_width_vs_elongation.svg"
        ).is_file()
        assert (Path(output["plot_directory"]) / "pca_interpretation.svg").is_file()
        assert (Path(output["plot_directory"]) / "pca_3d_feature_axes.svg").is_file()
        loaded = load_morphology_rows_from_workbook(output["workbook"])
        assert len(loaded) == 80
        workbook = load_workbook(output["workbook"], read_only=True, data_only=True)
        settings_rows = list(workbook["Morphology_Cluster_Settings"].iter_rows(values_only=True))
        equation_rows = list(workbook["PCA_Equations"].iter_rows(values_only=True))
        preparation_rows = list(
            workbook["PCA_Feature_Preparation"].iter_rows(values_only=True)
        )
        workbook.close()
        assert any("plot_style" in tuple(str(value) for value in row) for row in settings_rows)
        assert any("expanded_equation" in row for row in equation_rows)
        assert any("scaling_center" in row for row in preparation_rows)
        standalone = run_morphology_clustering_from_workbook(
            output["workbook"],
            root / "standalone.xlsx",
            settings,
        )
        assert standalone["result"]["included_spine_count"] == 80
        assert Path(standalone["export"]["workbook"]).is_file()
