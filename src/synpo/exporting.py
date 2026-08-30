from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Callable

import numpy as np
from openpyxl import Workbook, load_workbook

from .measurements import (
    distribution_summary_rows,
    load_distribution_preview,
    load_measurement_result,
)


Progress = Callable[[str, int, int, str], None]
BIN_COLORS = (
    (35, 0, 75), (75, 3, 110), (112, 14, 117), (147, 37, 103), (177, 63, 82),
    (204, 93, 58), (224, 127, 36), (239, 167, 25), (246, 210, 42), (240, 249, 33),
)


def _cell(value: object) -> object:
    if isinstance(value, (list, tuple, dict)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return value


def _columns(rows: list[dict[str, object]]) -> list[str]:
    columns: list[str] = []
    for row in rows:
        for key in row:
            if key not in columns:
                columns.append(key)
    return columns


def _group_summary(specimen_rows: list[dict[str, object]]) -> list[dict[str, object]]:
    output: list[dict[str, object]] = []
    groups = sorted({str(row.get("experimental_group", "")) for row in specimen_rows})
    excluded = {"experimental_group", "specimen_id", "average_protein_distribution"}
    numeric = sorted(
        {
            key
            for row in specimen_rows
            for key, value in row.items()
            if key not in excluded and isinstance(value, (int, float)) and not isinstance(value, bool)
        }
    )
    for group in groups:
        members = [row for row in specimen_rows if str(row.get("experimental_group", "")) == group]
        for metric in numeric:
            values = [float(row[metric]) for row in members if row.get(metric) is not None]
            if not values:
                continue
            sd = float(np.std(values, ddof=1)) if len(values) > 1 else 0.0
            output.append(
                {
                    "experimental_group": group,
                    "metric": metric,
                    "n_specimens": len(values),
                    "mean": float(np.mean(values)),
                    "sd": sd,
                    "sem": sd / np.sqrt(len(values)),
                }
            )
    return output


def collect_export_tables(manifest: dict[str, object]) -> dict[str, list[dict[str, object]]]:
    results = [
        load_measurement_result(manifest, index)
        for index, specimen in enumerate(manifest["specimens"])
        if specimen["checkpoints"].get("measurements", {}).get("state") == "complete"
    ]
    specimen = [row for result in results for row in result.get("specimen_rows", [])]
    dendrite = [row for result in results for row in result.get("dendrite_rows", [])]
    spine = [row for result in results for row in result.get("spine_rows", [])]
    clusters = [
        row
        for result in results
        for row in result.get("cluster_rows", [])
        if row.get("row_type") == "individual_cluster"
    ]
    sums = [
        row
        for result in results
        for row in result.get("cluster_rows", [])
        if row.get("row_type") == "spine_cluster_sum"
    ]
    distributions = [row for result in results for row in result.get("distribution_rows", [])]
    distribution_specimen, distribution_group = distribution_summary_rows(results)
    excluded = [
        row
        for row in distributions
        if bool(row.get("spine_valid", True)) and not bool(row.get("distribution_included", False))
    ]
    invalid_ids = {
        (str(row.get("experimental_group", "")), str(row.get("specimen_id", "")), int(row["spine_id"]))
        for row in spine
        if not bool(row.get("spine_valid", True))
    }
    invalid = [
        row
        for row in spine
        if (str(row.get("experimental_group", "")), str(row.get("specimen_id", "")), int(row["spine_id"])) in invalid_ids
    ]
    settings = [
        {"section": "calibration", **dict(manifest["calibration"])},
        {"section": "measurements", **dict(manifest["measurements"]["settings"])},
        {
            "section": "distribution",
            "bin_count": 10,
            "orientation": "largest shaft contact to longest curved distal skeleton path",
            "aggregation": "spine means within specimen, then specimen means within group",
            "group_error_bars": "SEM across specimen means",
            "ratio_units": "fraction 0-1",
        },
    ]
    return {
        "Specimen_Master": sorted(specimen, key=lambda row: (str(row.get("experimental_group")), str(row.get("specimen_id")))),
        "Dendrite_Master": sorted(dendrite, key=lambda row: (str(row.get("experimental_group")), str(row.get("specimen_id")), int(row.get("dendrite_id", 0)))),
        "Spine_Master": sorted(spine, key=lambda row: (str(row.get("experimental_group")), str(row.get("specimen_id")), int(row.get("spine_id", 0)))),
        "Cluster_Individual": clusters,
        "Cluster_Sums": sums,
        "Distribution_Individual": distributions,
        "Distribution_Specimen": distribution_specimen,
        "Distribution_Group": distribution_group,
        "Distribution_Excluded": excluded,
        "Invalid_Spines": invalid,
        "Group_Summary": _group_summary(specimen),
        "Settings": settings,
    }


def _write_workbook(path: Path, tables: dict[str, list[dict[str, object]]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in tables.items():
        sheet = workbook.create_sheet(name[:31])
        columns = _columns(rows)
        if columns:
            sheet.append(columns)
            for row in rows:
                sheet.append([_cell(row.get(column)) for column in columns])
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
        else:
            sheet.append(["No rows"])
    workbook.save(path)


def _write_csvs(directory: Path, tables: dict[str, list[dict[str, object]]]) -> list[Path]:
    directory.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    for name, rows in tables.items():
        path = directory / f"{name}.csv"
        columns = _columns(rows)
        with path.open("w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.DictWriter(stream, fieldnames=columns, extrasaction="ignore")
            if columns:
                writer.writeheader()
                for row in rows:
                    writer.writerow({key: _cell(row.get(key)) for key in columns})
        written.append(path)
    return written


def _pdf_pages(
    manifest: dict[str, object],
    path: Path,
    selected: list[tuple[int, dict[str, object]]],
    group_rows: list[dict[str, object]],
    margin_um: float,
    *,
    include_group_summary: bool,
    progress: Progress | None,
) -> None:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages

    colors = np.asarray(BIN_COLORS, dtype=np.float64) / 255.0
    with PdfPages(path) as pdf:
        wrote_page = False
        if include_group_summary:
            for group in group_rows:
                means = np.asarray(
                    [
                        np.nan if group.get(f"bin_{index:02d}_mean") is None else float(group[f"bin_{index:02d}_mean"])
                        for index in range(1, 11)
                    ],
                    dtype=np.float64,
                )
                sem = np.asarray(
                    [
                        np.nan if group.get(f"bin_{index:02d}_sem") is None else float(group[f"bin_{index:02d}_sem"])
                        for index in range(1, 11)
                    ],
                    dtype=np.float64,
                )
                figure, axis = plt.subplots(figsize=(9, 6))
                axis.errorbar(range(1, 11), means, yerr=sem, marker="o", capsize=3)
                axis.set(xlabel="Spine part (shaft → tip)", ylabel="Cluster volume / spine-part volume", title=f"{group['experimental_group']} — specimen-weighted mean ± SEM")
                axis.set_xlim(0.5, 10.5)
                axis.set_ylim(bottom=0)
                axis.grid(alpha=0.25)
                pdf.savefig(figure, bbox_inches="tight")
                plt.close(figure)
                wrote_page = True
        for position, (specimen_index, row) in enumerate(selected):
            preview = load_distribution_preview(
                manifest, specimen_index, int(row["spine_id"]), margin_um=margin_um
            )
            figure = plt.figure(figsize=(11.7, 8.3))
            grid = figure.add_gridspec(2, 2, height_ratios=(3, 2))
            dendrite_axis = figure.add_subplot(grid[0, 0])
            protein_axis = figure.add_subplot(grid[0, 1])
            profile_axis = figure.add_subplot(grid[1, :])
            for axis, raw, bins, title in (
                (dendrite_axis, preview.dendrite_projection, preview.spine_bins_projection, "Dendrite/spine channel"),
                (protein_axis, preview.protein_projection, preview.cluster_bins_projection, "Protein channel"),
            ):
                axis.imshow(raw, cmap="gray", vmin=np.percentile(raw, 0.5), vmax=np.percentile(raw, 99.8))
                overlay = np.zeros((*bins.shape, 4), dtype=np.float64)
                for bin_index in range(1, 11):
                    overlay[bins == bin_index] = (*colors[bin_index - 1, :3], 0.55 if axis is dendrite_axis else 0.85)
                axis.imshow(overlay)
                if np.any(bins):
                    axis.contour(bins > 0, levels=[0.5], colors="white", linewidths=0.7)
                axis.set_title(title)
                axis.axis("off")
            if preview.axis_xy:
                dendrite_axis.plot(
                    [point[0] for point in preview.axis_xy],
                    [point[1] for point in preview.axis_xy],
                    color="white",
                    linewidth=1.2,
                )
            if preview.base_point_local_zyx is not None:
                dendrite_axis.scatter(
                    [preview.base_point_local_zyx[2]],
                    [preview.base_point_local_zyx[1]],
                    s=42,
                    c="#20d060",
                    edgecolors="black",
                    linewidths=0.6,
                    label="automatic base",
                    zorder=5,
                )
            if preview.endpoint_local_zyx is not None:
                dendrite_axis.scatter(
                    [preview.endpoint_local_zyx[2]],
                    [preview.endpoint_local_zyx[1]],
                    s=42,
                    c="#ed32c8",
                    edgecolors="black",
                    linewidths=0.6,
                    label=f"{row.get('centerline_endpoint_source', 'automatic')} endpoint",
                    zorder=5,
                )
            if preview.base_point_local_zyx is not None or preview.endpoint_local_zyx is not None:
                dendrite_axis.legend(loc="lower right", fontsize=7)
            ratios = [row.get(f"bin_{index:02d}_ratio") for index in range(1, 11)]
            profile_axis.plot(range(1, 11), ratios, marker="o")
            profile_axis.set(xlabel="Spine part (shaft → tip)", ylabel="Cluster volume / spine-part volume")
            profile_axis.set_ylim(bottom=0)
            profile_axis.grid(alpha=0.25)
            figure.suptitle(
                f"{row['experimental_group']} | {row['specimen_id']} | spine {row['spine_id']}\n"
                f"Axis: {row['distribution_axis_status']} | included: {row.get('distribution_included')} | "
                f"valid: {row.get('spine_valid')} | endpoint: "
                f"{row.get('centerline_endpoint_source', 'automatic')} | {row.get('review_note', '')}"
            )
            pdf.savefig(figure, bbox_inches="tight")
            plt.close(figure)
            wrote_page = True
            if progress:
                progress("Exporting validation PDF", position + 1, len(selected), f"Spine {row['spine_id']}")
        if not wrote_page:
            figure, axis = plt.subplots(figsize=(9, 6))
            axis.axis("off")
            axis.text(
                0.5,
                0.5,
                "No spines matched this optional PDF category.",
                ha="center",
                va="center",
                fontsize=14,
            )
            pdf.savefig(figure, bbox_inches="tight")
            plt.close(figure)


def export_measurements(
    manifest: dict[str, object],
    workbook_path: str | Path,
    *,
    validation_pdf: bool = False,
    excluded_audit_pdf: bool = False,
    invalid_audit_pdf: bool = False,
    pdf_margin_um: float = 1.0,
    progress: Progress | None = None,
) -> dict[str, object]:
    path = Path(workbook_path).resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    tables = collect_export_tables(manifest)
    _write_workbook(path, tables)
    csv_directory = path.parent / f"{path.stem}_csv"
    csv_paths = _write_csvs(csv_directory, tables)
    written_pdfs: list[Path] = []
    indexed_rows: list[tuple[int, dict[str, object]]] = []
    for specimen_index, _specimen in enumerate(manifest["specimens"]):
        try:
            result = load_measurement_result(manifest, specimen_index)
        except ValueError:
            continue
        indexed_rows.extend((specimen_index, row) for row in result.get("distribution_rows", []))
    selections = []
    if validation_pdf:
        selections.append((path.with_name(f"{path.stem}_distribution_validation.pdf"), [item for item in indexed_rows if bool(item[1].get("spine_valid", True)) and bool(item[1].get("distribution_included", False))], True))
    if excluded_audit_pdf:
        selections.append((path.with_name(f"{path.stem}_distribution_excluded_audit.pdf"), [item for item in indexed_rows if bool(item[1].get("spine_valid", True)) and not bool(item[1].get("distribution_included", False))], False))
    if invalid_audit_pdf:
        selections.append((path.with_name(f"{path.stem}_invalid_spines_audit.pdf"), [item for item in indexed_rows if not bool(item[1].get("spine_valid", True))], False))
    for pdf_path, selected, summaries in selections:
        _pdf_pages(
            manifest,
            pdf_path,
            selected,
            tables["Distribution_Group"],
            pdf_margin_um,
            include_group_summary=summaries,
            progress=progress,
        )
        written_pdfs.append(pdf_path)

    # Verification happens before the UI may offer any future cache deletion.
    load_workbook(path, read_only=True).close()
    if not csv_paths or any(not csv_path.is_file() for csv_path in csv_paths):
        raise OSError("CSV export verification failed.")
    for pdf_path in written_pdfs:
        with pdf_path.open("rb") as stream:
            if stream.read(4) != b"%PDF":
                raise OSError(f"PDF verification failed: {pdf_path.name}")
    return {
        "workbook": str(path),
        "csv_directory": str(csv_directory),
        "pdfs": [str(item) for item in written_pdfs],
        "verified": True,
    }
